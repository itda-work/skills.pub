"""값 grid 로드 — openpyxl 계산값(data_only) / CSV. #967 data-verify.

백엔드 독립: 이 모듈은 (sheet_name -> Grid) 만 공급하고, 판정은 verifiers 가 한다.
다른 로더가 같은 Grid 인터페이스를 채우면 verifiers 를 그대로 재사용할 수 있다.

입력 방어(itda-work/skills#30):
  · grid 는 **값이 있는 범위**까지만 만든다. 예전 `ws.iter_rows()` 는 서식만 남은 먼 칸까지 빈 칸마다
    Cell 을 만들어, 서식 하나(200000행×300열)로 90초 넘게 멈췄다(1만×100 = 1.4초·220MB 실측).
    빈 행은 한 리스트를 함께 가리켜 행 수가 커도 메모리가 칸 수만큼 늘지 않는다.
  · 병합 범위 방어는 `xlsx_guard` 가 openpyxl 보다 먼저 한다 — 한 일은 Note 로 돌려준다.
  · date1904 워크북의 날짜 보정은 openpyxl 이 한다(`wb.epoch`) — 테스트가 고정한다.
"""
from __future__ import annotations

import csv
from dataclasses import dataclass, field

try:
    import openpyxl
    from openpyxl.utils.cell import coordinate_to_tuple
except ImportError as exc:  # 명시적 에러(조용한 폴백 금지)
    raise SystemExit("openpyxl 필요 — python3 -m pip install -r requirements.txt") from exc

import xlsx_guard


@dataclass
class Grid:
    name: str
    rows: list = field(default_factory=list)   # list[list] 계산값(헤더 포함)

    def header(self) -> list:
        return self.rows[0] if self.rows else []

    def data_rows(self) -> list:
        return self.rows[1:] if self.rows else []

    def col_index(self, name) -> int | None:
        for i, v in enumerate(self.header()):
            if str(v).strip() == str(name).strip():
                return i
        return None

    def cell_a1(self, ref: str):
        r, c = coordinate_to_tuple(ref)   # 1-based (row, col)
        if 1 <= r <= len(self.rows) and 1 <= c <= len(self.rows[r - 1]):
            return self.rows[r - 1][c - 1]
        return None


def _load_csv(path: str) -> Grid:
    for enc in ("utf-8-sig", "utf-8", "cp949"):
        try:
            with open(path, newline="", encoding=enc) as f:
                return Grid(name="Sheet1", rows=[list(r) for r in csv.reader(f)])
        except UnicodeDecodeError:
            continue
    raise ValueError(f"CSV 인코딩 판별 실패(utf-8/cp949): {path}")


def _grid_from_ws(name: str, ws) -> Grid:
    """파일에 실제로 있는 칸(openpyxl 비공개 `_cells`)만 읽어 값 범위의 grid 를 만든다."""
    cells = getattr(ws, "_cells", None)
    if not isinstance(cells, dict):
        raise RuntimeError("openpyxl 워크시트에 _cells 가 없습니다 — openpyxl 버전을 확인하세요(>=3.1)")
    values = {rc: c.value for rc, c in cells.items() if c.value is not None}
    if not values:
        return Grid(name=name, rows=[])
    n_rows = max(r for r, _ in values)
    n_cols = max(c for _, c in values)
    empty = [None] * n_cols            # 값 없는 행은 이 한 리스트를 함께 가리킨다(읽기 전용)
    by_row: dict[int, list] = {}
    for (r, c), v in values.items():
        row = by_row.get(r)
        if row is None:
            row = by_row[r] = [None] * n_cols
        row[c - 1] = v
    return Grid(name=name, rows=[by_row.get(r, empty) for r in range(1, n_rows + 1)])


def load_sheets_with_notes(path: str) -> tuple[dict, list]:
    """({sheet_name: Grid}, [xlsx_guard.Note]). .csv/.tsv 는 단일 Sheet1·Note 없음."""
    if path.lower().endswith((".csv", ".tsv")):
        return {"Sheet1": _load_csv(path)}, []
    with xlsx_guard.prepared(path) as (load_path, notes):
        wb = openpyxl.load_workbook(load_path, data_only=True)
    return {name: _grid_from_ws(name, wb[name]) for name in wb.sheetnames}, notes


def load_sheets(path: str) -> dict:
    """{sheet_name: Grid}. .xlsx 는 전체 시트(계산값), .csv/.tsv 는 단일 Sheet1."""
    return load_sheets_with_notes(path)[0]
