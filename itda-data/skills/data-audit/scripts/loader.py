"""워크북 로드 — openpyxl 수식면·값면 양면 (#952 audit-xls 이식).

audit 은 '누가 =A1*1.05 를 박았나'(수식 문자열)와 '#REF! 가 떴나'(캐시 계산값) 양쪽이 필요하다.
openpyxl 은 한 번의 load 로 둘 다 주지 않는다 — data_only=False 는 수식, True 는 마지막 저장 시 캐시값.
그래서 두 번 로드해 시트별 formulas/computed/literals 로 가른다.

입력 방어(itda-work/skills#30):
  · 칸 순회는 `ws._cells`(파일에 실제로 있는 칸)만 돈다. `iter_rows(max_row, max_col)` 는 빈 칸마다 Cell 을
    새로 만들어, 먼 칸에 서식 하나만 남은 시트(200000행×300열)에서 90초 넘게 멈췄다(1만×100 = 6.6초·447MB 실측).
  · 병합 범위 방어는 `xlsx_guard` 가 openpyxl 보다 먼저 한다(거꾸로 된·거대한 병합, 가려진 값).
  · date1904 워크북의 날짜 보정은 openpyxl 이 한다(`wb.epoch`) — 테스트가 고정한다.

한계: data_only=True 캐시값은 '파일에 마지막 저장된' 값이다. openpyxl 로만 만들고 Excel 로
저장한 적 없는 파일은 캐시가 비어 computed=None 이 된다(SKILL.md 정확성 한계 절 참조).
"""
from __future__ import annotations

from dataclasses import dataclass, field

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # 명시적 에러 표면화(조용한 폴백 금지)
    raise SystemExit(
        "openpyxl 이 필요합니다 — python3 -m pip install -r requirements.txt"
    ) from exc

import xlsx_guard

ERROR_CODES = ("#REF!", "#VALUE!", "#N/A", "#DIV/0!", "#NAME?", "#NULL!", "#NUM!")


def a1(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"


def is_error_value(val) -> bool:
    return isinstance(val, str) and val in ERROR_CODES


@dataclass
class SheetView:
    name: str
    state: str                 # 'visible' | 'hidden' | 'veryHidden'
    max_row: int
    max_col: int
    formulas: dict             # (row, col) -> 수식 문자열 '=...'
    computed: dict             # (row, col) -> 수식의 캐시 계산값(에러 문자열 포함, None 가능)
    literals: dict             # (row, col) -> 정적 값(수식 아님)
    hidden_rows: list          # [int, ...]
    hidden_cols: list          # ['A', ...]
    input_notes: list = field(default_factory=list)   # xlsx_guard.Note — 입력 방어로 한 일


def _existing_cells(ws):
    """파일에 실제로 있는 칸만 (row, col) 순으로. openpyxl 비공개 `_cells` 를 쓴다 — 공개 API 는
    범위 순회뿐이고 그것은 빈 칸마다 객체를 만든다. `_cells` 가 사라지면 명시적으로 실패한다."""
    cells = getattr(ws, "_cells", None)
    if not isinstance(cells, dict):
        raise RuntimeError("openpyxl 워크시트에 _cells 가 없습니다 — openpyxl 버전을 확인하세요(>=3.1)")
    return sorted(cells.items())


def load_views(path: str, sheet: str | None = None):
    """(list[SheetView], list[all_sheet_names]) 반환. sheet=None 이면 전체(숨은 시트 포함)."""
    with xlsx_guard.prepared(path) as (load_path, notes):
        wb_f = openpyxl.load_workbook(load_path, data_only=False)
        wb_v = openpyxl.load_workbook(load_path, data_only=True)
    all_names = list(wb_f.sheetnames)
    if sheet is not None and sheet not in all_names:
        raise ValueError(f"시트를 찾을 수 없습니다: {sheet!r} (가능: {all_names})")
    targets = all_names if sheet is None else [sheet]

    views: list[SheetView] = []
    for name in targets:
        ws_f, ws_v = wb_f[name], wb_v[name]
        cells_v = getattr(ws_v, "_cells", {})
        formulas: dict = {}
        computed: dict = {}
        literals: dict = {}
        max_row = max_col = 0
        for rc, cf in _existing_cells(ws_f):
            v = cf.value
            if cf.data_type == "f" or (isinstance(v, str) and v.startswith("=")):
                formulas[rc] = v
                cv = cells_v.get(rc)
                computed[rc] = cv.value if cv is not None else None
            elif v is not None:
                literals[rc] = v
            else:
                continue
            max_row, max_col = max(max_row, rc[0]), max(max_col, rc[1])
        hidden_rows = sorted(i for i, d in ws_f.row_dimensions.items() if d.hidden)
        hidden_cols = sorted(k for k, d in ws_f.column_dimensions.items() if d.hidden)
        views.append(
            SheetView(
                name=name,
                state=ws_f.sheet_state,
                max_row=max_row,
                max_col=max_col,
                formulas=formulas,
                computed=computed,
                literals=literals,
                hidden_rows=hidden_rows,
                hidden_cols=hidden_cols,
                input_notes=[n for n in notes if n.sheet == name],
            )
        )
    return views, all_names
